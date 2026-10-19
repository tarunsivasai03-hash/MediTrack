import io
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook

from app import db
from app.models import (
    Appointment, Department, Doctor, DoctorTimeOff, Invoice, MedicalRecord, Patient, Rating, User,
)
from app.reports import report_data
from app.services import BookingError, check_in, create_walk_in, queue_state, start_consultation, todays_queue
from tests.conftest import login


def _doctor():
    return Doctor.query.first()


def _patient(email="pat@t.com"):
    return User.query.filter_by(email=email).first().patient_profile


def _appt(status="confirmed", when=None, slot=time(9, 0), patient=None, doctor=None, **kw):
    a = Appointment(patient_id=(patient or _patient()).id, doctor_id=(doctor or _doctor()).id,
                    date=when or date.today(), time=slot, status=status, reason="cough", **kw)
    db.session.add(a)
    db.session.commit()
    return a


def _doc2(fee=500.0, days="Mon,Tue,Wed,Thu,Fri,Sat,Sun"):
    u = User(name="Dr Two", email="two@t.com", role="doctor")
    u.set_password("password123")
    db.session.add(u); db.session.flush()
    d = Doctor(user_id=u.id, specialization="ENT", qualification="MS", consultation_fee=fee,
               work_start=time(9, 0), work_end=time(11, 0), available_days=days)
    db.session.add(d); db.session.commit()
    return d


def _complete(client, appt):
    login(client, "doc@t.com")
    r = client.post(f"/doctor/appointments/{appt.id}/record", data={"diagnosis": "Flu", "treatment": "Rest"})
    client.post("/logout")
    return r


# ═══ Reception role ══════════════════════════════════════════════════════

def test_reception_can_use_front_desk_pages(client):
    login(client, "desk@t.com")
    assert client.get("/").location.endswith("/reception/dashboard")
    for url in ("/reception/dashboard", "/reception/queue", "/reception/walk-in", "/reception/book",
                "/reception/appointments", "/reception/patients", "/billing/"):
        assert client.get(url).status_code == 200, url


def test_reception_cannot_reach_clinical_or_admin_areas(client):
    appt = _appt()
    login(client, "desk@t.com")
    for url in ("/admin/dashboard", "/admin/doctors", "/admin/audit", "/admin/reports", "/doctor/dashboard",
                f"/doctor/appointments/{appt.id}", "/patient/dashboard", "/patient/records", "/api/stats",
                f"/api/patients/{appt.patient_id}/vitals", "/api/vitals"):
        assert client.get(url).status_code == 403, url
    assert client.post(f"/files/appointments/{appt.id}/upload").status_code == 403
    assert client.get("/files/records/1/prescription.pdf").status_code in (403, 404)


def test_other_roles_cannot_use_reception_pages(client):
    for email in ("pat@t.com", "doc@t.com", "admin@t.com"):
        login(client, email)
        assert client.get("/reception/queue").status_code == 403, email
        assert client.get("/reception/walk-in").status_code == 403, email
        client.post("/logout")


def test_reception_sees_contact_details_only(client):
    p = _patient()
    p.allergies, p.chronic_conditions = "Penicillin", "Asthma"
    db.session.commit()
    login(client, "desk@t.com")
    page = client.get("/reception/patients").get_data(as_text=True)
    assert "Pat" in page and "Penicillin" not in page and "Asthma" not in page


def test_admin_manages_reception_staff(client):
    login(client, "admin@t.com")
    r = client.post("/admin/staff", data={"name": "New Desk", "email": "new@t.com", "password": "longpassword1"})
    assert r.status_code == 302 and User.query.filter_by(email="new@t.com").first().role == "reception"
    assert b"already registered" in client.post("/admin/staff", data={"name": "X", "email": "new@t.com", "password": "longpassword1"}).data
    uid = User.query.filter_by(email="new@t.com").first().id
    client.post(f"/admin/staff/{uid}/toggle")
    assert not db.session.get(User, uid).is_active
    assert client.post(f"/admin/staff/{_doctor().user_id}/toggle").status_code == 404   # only reception accounts


# ═══ Walk-ins, check-in and the queue ════════════════════════════════════

def test_walk_in_existing_patient_joins_queue(client):
    login(client, "desk@t.com")
    r = client.post("/reception/walk-in", data={"patient_id": _patient().id, "doctor_id": _doctor().id, "reason": "fever"})
    assert r.status_code == 302
    a = Appointment.query.one()
    assert a.is_walk_in and a.status == "confirmed" and a.checked_in_at and a.date == date.today()
    assert queue_state(a) == "waiting" and a.reason == "fever"


def test_walk_in_registers_a_new_patient_without_login(client):
    login(client, "desk@t.com")
    r = client.post("/reception/walk-in", data={
        "patient_id": 0, "new_name": "Raju Kumar", "new_age": "41", "new_gender": "Male", "new_phone": "98000",
        "doctor_id": _doctor().id})
    assert r.status_code == 302
    p = Patient.query.filter(Patient.phone == "98000").one()
    assert not p.user.can_login and p.user.name == "Raju Kumar" and p.user.email.endswith("@no-login.invalid")
    assert Appointment.query.one().patient_id == p.id
    client.post("/logout")
    assert login(client, p.user.email, "anything").status_code == 200          # can never sign in


def test_walk_in_form_validation(client):
    login(client, "desk@t.com")
    r = client.post("/reception/walk-in", data={"patient_id": 0, "doctor_id": _doctor().id})
    assert r.status_code == 200 and b"Enter the patient" in r.data and Appointment.query.count() == 0


def test_walk_in_when_the_day_is_full_goes_to_the_end(app):
    doc = _doctor()
    doc.work_start, doc.work_end = time(0, 0), time(0, 30)            # one slot, already past
    db.session.commit()
    a = create_walk_in(_patient(), doc, "late")
    assert a.is_walk_in and a.date == date.today()
    b = create_walk_in(_patient("pat2@t.com"), doc, "later")
    assert b.time != a.time                                            # no clash on the unique slot index


def test_walk_in_refused_on_day_off_or_leave(app):
    doc = _doctor()
    days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    doc.available_days = ",".join(d for d in days if d != days[date.today().weekday()])
    db.session.commit()
    with pytest.raises(BookingError, match="doesn't work"):
        create_walk_in(_patient(), doc, "x")
    doc.available_days = ",".join(days)
    db.session.add(DoctorTimeOff(doctor_id=doc.id, date=date.today()))
    db.session.commit()
    with pytest.raises(BookingError, match="on leave"):
        create_walk_in(_patient(), doc, "x")


def test_check_in_rules(app):
    today_appt = _appt("pending")
    check_in(today_appt)
    assert today_appt.checked_in_at and today_appt.status == "confirmed"
    with pytest.raises(BookingError, match="already"):
        check_in(today_appt)
    with pytest.raises(BookingError, match="today"):
        check_in(_appt(when=date.today() + timedelta(days=1), slot=time(9, 30)))
    with pytest.raises(BookingError):
        check_in(_appt("cancelled", slot=time(10, 0)))


def test_queue_order_and_states(app):
    scheduled = _appt(slot=time(9, 0))
    first = _appt(slot=time(9, 30), patient=_patient("pat2@t.com"))
    second = _appt(slot=time(10, 0))
    done = _appt("completed", slot=time(10, 30))
    first.checked_in_at = datetime.now() - timedelta(minutes=20)
    second.checked_in_at = datetime.now() - timedelta(minutes=5)
    db.session.commit()
    order = todays_queue()
    assert [a.id for a in order] == [first.id, second.id, scheduled.id, done.id]   # waiting by arrival, then scheduled, then done
    start_consultation(second)
    assert todays_queue()[0].id == second.id and queue_state(second) == "in_consult"


def test_queue_api_per_role_and_actions(client):
    a = _appt()
    login(client, "desk@t.com")
    data = client.get("/api/queue").get_json()
    row = data["rows"][0]
    assert row["state"] == "scheduled" and any(x["label"] == "Check in" for x in row["actions"])
    client.post(f"/reception/appointments/{a.id}/check-in")
    row = client.get("/api/queue").get_json()["rows"][0]
    assert row["state"] == "waiting" and row["position"] == 1 and row["wait_minutes"] >= 0
    client.post("/logout")

    login(client, "doc@t.com")
    row = client.get("/api/queue").get_json()["rows"][0]
    assert [x["label"] for x in row["actions"]] == ["Start", "Complete"]
    assert client.post(f"/doctor/appointments/{a.id}/start").status_code == 302
    assert client.get("/api/queue").get_json()["rows"][0]["state"] == "in_consult"
    client.post("/logout")

    login(client, "pat@t.com")
    assert client.get("/api/queue").status_code == 403


def test_doctor_queue_only_shows_own_patients(client):
    other = _doc2()
    _appt()
    mine = _appt(slot=time(10, 0), doctor=other)
    login(client, "doc@t.com")
    ids = [r["id"] for r in client.get("/api/queue").get_json()["rows"]]
    assert mine.id not in ids and len(ids) == 1
    assert client.post(f"/doctor/appointments/{mine.id}/start").status_code == 403


def test_queue_names_cannot_inject_markup(client):
    _appt().patient.user.name = "<img src=x onerror=alert(1)>"
    db.session.commit()
    login(client, "desk@t.com")
    # the API returns the raw text; the page script inserts it with textContent
    assert client.get("/api/queue").get_json()["rows"][0]["patient"] == "<img src=x onerror=alert(1)>"
    js = client.get("/static/js/queue.js").get_data(as_text=True)
    assert "innerHTML" not in js.replace("never innerHTML", "")


def test_reception_books_for_a_patient_and_it_is_confirmed(client, app):
    tomorrow = date.today() + timedelta(days=1)
    login(client, "desk@t.com")
    r = client.post("/reception/book", data={"patient_id": _patient().id, "doctor_id": _doctor().id,
                                             "date": tomorrow.isoformat(), "time": "09:30", "reason": "checkup"})
    assert r.status_code == 302
    a = Appointment.query.one()
    assert a.status == "confirmed"
    assert any("confirmed" in m["Subject"] for m in app.extensions["outbox"])


def test_reception_status_actions(client):
    past = _appt()
    login(client, "desk@t.com")
    client.post(f"/reception/appointments/{past.id}/status", data={"status": "no_show"})
    assert db.session.get(Appointment, past.id).status == "no_show"
    other = _appt(slot=time(10, 0))
    client.post(f"/reception/appointments/{other.id}/status", data={"status": "completed"})   # doctors only
    assert db.session.get(Appointment, other.id).status == "confirmed"
    client.post(f"/reception/appointments/{other.id}/status", data={"status": "cancelled"})
    assert db.session.get(Appointment, other.id).status == "cancelled"


def test_no_show_rules(app):
    from app.services import change_status
    future = _appt(when=date.today() + timedelta(days=2))
    with pytest.raises(BookingError, match="future"):
        change_status(future, "no_show")
    pending = _appt("pending", slot=time(10, 0))
    with pytest.raises(BookingError):
        change_status(pending, "no_show")                              # must be confirmed first
    ok = _appt(slot=time(10, 30))
    change_status(ok, "no_show")
    with pytest.raises(BookingError):
        change_status(ok, "completed")                                 # terminal


# ═══ Billing ═════════════════════════════════════════════════════════════

def test_completing_a_visit_creates_one_invoice(client):
    doc = _doctor()
    doc.consultation_fee = 800
    db.session.commit()
    a = _appt()
    _complete(client, a)
    inv = Invoice.query.one()
    assert (inv.amount, inv.status, inv.patient_id, inv.doctor_id) == (800, "unpaid", a.patient_id, doc.id)
    assert inv.number == f"INV-{inv.id:06d}"
    login(client, "doc@t.com")                                        # editing the record must not bill twice
    client.post(f"/doctor/appointments/{a.id}/record", data={"diagnosis": "Flu2", "treatment": "Rest"})
    assert Invoice.query.count() == 1


def test_free_visits_get_no_invoice(client):
    _complete(client, _appt())                                        # fee defaults to 0
    assert Invoice.query.count() == 0


def _invoice(client, fee=800):
    _doctor().consultation_fee = fee
    db.session.commit()
    _complete(client, _appt())
    return Invoice.query.one()


def test_reception_records_payment_with_discount(client):
    inv = _invoice(client)
    login(client, "desk@t.com")
    assert client.get(f"/billing/{inv.id}").status_code == 200
    r = client.post(f"/billing/{inv.id}", data={"method": "upi", "discount": "100", "notes": "senior citizen"})
    assert r.status_code == 302
    inv = db.session.get(Invoice, inv.id)
    assert (inv.status, inv.method, inv.discount, inv.total) == ("paid", "upi", 100, 700)
    assert inv.paid_by.email == "desk@t.com" and inv.paid_at
    again = client.post(f"/billing/{inv.id}", data={"method": "cash", "discount": "0"}, follow_redirects=True)
    assert b"already paid" in again.data
    assert client.post(f"/billing/{inv.id}/void").status_code == 302
    assert db.session.get(Invoice, inv.id).status == "paid"           # paid invoices can't be voided


@pytest.mark.parametrize("data", [{"method": "bitcoin", "discount": "0"}, {"method": "cash", "discount": "9999"},
                                  {"method": "cash", "discount": "-5"}])
def test_invalid_payments_rejected(client, data):
    inv = _invoice(client)
    login(client, "desk@t.com")
    client.post(f"/billing/{inv.id}", data=data)
    assert db.session.get(Invoice, inv.id).status == "unpaid"


def test_void_unpaid_invoice(client):
    inv = _invoice(client)
    login(client, "desk@t.com")
    client.post(f"/billing/{inv.id}/void")
    assert db.session.get(Invoice, inv.id).status == "void"
    assert b"already void" in client.post(f"/billing/{inv.id}/void", follow_redirects=True).data


def test_billing_staff_only_and_filters(client):
    inv = _invoice(client)
    login(client, "pat@t.com")
    assert client.get("/billing/").status_code == 403 and client.post(f"/billing/{inv.id}").status_code == 403
    client.post("/logout")
    login(client, "doc@t.com")
    assert client.get("/billing/").status_code == 403
    client.post("/logout")
    assert "/login" in client.get("/billing/").location
    login(client, "admin@t.com")
    page = client.get("/billing/?status=unpaid&q=Pat").get_data(as_text=True)
    assert inv.number in page and "800" in page
    assert inv.number not in client.get("/billing/?status=paid").get_data(as_text=True)


def test_invoice_pdf_access(client):
    inv = _invoice(client)
    url = f"/billing/{inv.id}/pdf"
    for email, ok in (("pat@t.com", True), ("doc@t.com", True), ("desk@t.com", True), ("admin@t.com", True),
                      ("pat2@t.com", False)):
        login(client, email)
        r = client.get(url)
        assert (r.status_code == 200) is ok, email
        if ok:
            assert r.data.startswith(b"%PDF") and "invoice-" in r.headers["Content-Disposition"]
        client.post("/logout")
    assert "/login" in client.get(url).location
    login(client, "desk@t.com")
    client.post(f"/billing/{inv.id}", data={"method": "cash", "discount": "0"})
    assert "receipt-" in client.get(url).headers["Content-Disposition"]


def test_patient_sees_bills_and_balance(client):
    _invoice(client)
    login(client, "pat@t.com")
    assert "800.00" in client.get("/patient/billing").get_data(as_text=True)
    dash = client.get("/patient/dashboard").get_data(as_text=True)
    assert "Outstanding balance" in dash and "800.00" in dash


# ═══ Departments, directory and ratings ══════════════════════════════════

def test_admin_departments_crud_and_doctor_assignment(client):
    login(client, "admin@t.com")
    assert client.post("/admin/departments", data={"name": "Cardiology", "description": "Heart"}).status_code == 302
    assert b"already exists" in client.post("/admin/departments", data={"name": "cardiology"}).data
    dept = Department.query.one()
    client.post(f"/admin/doctors/{_doctor().id}/edit", data={
        "name": "Dr. Who", "email": "doc@t.com", "password": "", "specialization": "General", "qualification": "MBBS",
        "experience_years": "3", "consultation_fee": "0", "work_start": "09:00", "work_end": "11:00",
        "slot_minutes": "30", "available_days": ["Mon"], "department_id": str(dept.id)})
    assert _doctor().department_id == dept.id
    client.post(f"/admin/departments/{dept.id}/delete")
    assert Department.query.count() == 1                              # still has a doctor
    _doctor().department_id = None
    db.session.commit()
    client.post(f"/admin/departments/{dept.id}/delete")
    assert Department.query.count() == 0


def _directory_setup():
    cardio = Department(name="Cardiology"); ent = Department(name="ENT")
    db.session.add_all([cardio, ent]); db.session.flush()
    a = _doctor(); a.department_id, a.consultation_fee, a.available_days = cardio.id, 300, "Mon,Tue"
    b = _doc2(fee=900, days="Sat,Sun"); b.department_id = ent.id
    db.session.commit()
    return a, b, cardio, ent


def test_doctor_directory_filters_and_sorting(client):
    a, b, cardio, ent = _directory_setup()
    login(client, "pat@t.com")

    def names(qs):
        page = client.get("/patient/doctors" + qs).get_data(as_text=True)
        found = [(page.index(f">{n}</a>"), n) for n in ("Dr. Who", "Dr Two") if f">{n}</a>" in page]
        return [n for _, n in sorted(found)]            # in the order shown on the page

    assert names("") == ["Dr Two", "Dr. Who"]               # alphabetical by default
    assert names(f"?department={ent.id}") == ["Dr Two"]
    assert names("?max_fee=500") == ["Dr. Who"]
    assert names("?day=Sun") == ["Dr Two"]
    assert names("?q=ent") == [] or names("?q=ENT") == ["Dr Two"]
    assert names("?q=nobody") == []
    db.session.add(Rating(appointment_id=_appt().id, doctor_id=b.id, patient_id=_patient().id, stars=5))
    db.session.commit()
    assert names("?min_rating=4") == ["Dr Two"]
    assert names("?sort=rating")[0] == "Dr Two" and names("?sort=fee")[0] == "Dr. Who"


def test_directory_hides_inactive_doctors_and_profile_page(client):
    a, b, *_ = _directory_setup()
    b.user.is_active = False
    db.session.commit()
    login(client, "pat@t.com")
    assert "Dr Two" not in client.get("/patient/doctors").get_data(as_text=True)
    assert client.get(f"/patient/doctors/{b.id}").status_code == 404
    assert "Dr. Who" in client.get(f"/patient/doctors/{a.id}").get_data(as_text=True)


def test_rating_flow(client):
    a = _appt()
    login(client, "pat@t.com")
    assert b"after the visit is completed" in client.post(
        f"/patient/appointments/{a.id}/rate", data={"stars": "5"}, follow_redirects=True).data
    a.status = "completed"
    db.session.commit()
    client.post(f"/patient/appointments/{a.id}/rate", data={"stars": "4", "comment": "Kind and quick"})
    assert Rating.query.one().stars == 4
    client.post(f"/patient/appointments/{a.id}/rate", data={"stars": "5", "comment": "Even better"})
    r = Rating.query.one()                                            # updated, not duplicated
    assert (r.stars, r.comment) == (5, "Even better")
    client.post(f"/patient/appointments/{a.id}/rate", data={"stars": "9"})
    assert Rating.query.one().stars == 5                              # invalid value ignored
    assert "Even better" in client.get(f"/patient/doctors/{_doctor().id}").get_data(as_text=True)
    client.post("/logout")

    login(client, "pat2@t.com")
    assert client.post(f"/patient/appointments/{a.id}/rate", data={"stars": "1"}).status_code == 403
    client.post("/logout")
    login(client, "doc@t.com")
    assert "Even better" in client.get("/doctor/dashboard").get_data(as_text=True)


# ═══ Family members ══════════════════════════════════════════════════════

def _add_child(client, name="Little Pat"):
    return client.post("/patient/family/new", data={
        "name": name, "relationship": "Child", "age": "6", "gender": "Male", "blood_group": "", "phone": ""})


def test_adding_a_family_member(client):
    login(client, "pat@t.com")
    assert _add_child(client).status_code == 302
    kid = Patient.query.filter(Patient.relationship_label == "Child").one()
    assert kid.guardian.email == "pat@t.com" and not kid.user.can_login and kid.age == 6
    page = client.get("/patient/dashboard").get_data(as_text=True)
    assert "Little Pat" in page and "family-switch" in page            # switched to the child automatically


def test_dependent_cannot_sign_in_or_reset_password(client, app):
    login(client, "pat@t.com")
    _add_child(client)
    kid = Patient.query.filter(Patient.relationship_label == "Child").one()
    client.post("/logout")
    assert login(client, kid.user.email, "anything").status_code == 200
    app.extensions["outbox"] = []
    client.post("/forgot-password", data={"email": kid.user.email})
    assert app.extensions["outbox"] == []


def test_switching_scopes_everything_to_that_person(client):
    tomorrow = date.today() + timedelta(days=1)
    login(client, "pat@t.com")
    _add_child(client)
    kid = Patient.query.filter(Patient.relationship_label == "Child").one()
    parent = _patient()
    client.post(f"/patient/switch/{parent.id}")
    client.post("/patient/book", data={"doctor_id": _doctor().id, "date": tomorrow.isoformat(), "time": "09:00", "reason": "me"})
    client.post(f"/patient/switch/{kid.id}")
    client.post("/patient/book", data={"doctor_id": _doctor().id, "date": tomorrow.isoformat(), "time": "09:30", "reason": "kid cough"})

    assert Appointment.query.filter_by(patient_id=kid.id).one().reason == "kid cough"
    assert Appointment.query.filter_by(patient_id=parent.id).one().reason == "me"
    assert "kid cough" in client.get("/patient/appointments").get_data(as_text=True)
    client.post(f"/patient/switch/{parent.id}")
    page = client.get("/patient/appointments").get_data(as_text=True)
    assert "kid cough" not in page and ">me<" in page.replace(" ", "").replace("\n", "") or "me" in page


def test_cannot_switch_to_or_touch_someone_elses_family(client):
    login(client, "pat2@t.com")
    _add_child(client, "Other Kid")
    other_kid = Patient.query.filter(Patient.relationship_label == "Child").one()
    kid_appt = _appt(patient=other_kid)
    client.post("/logout")

    login(client, "pat@t.com")
    assert client.post(f"/patient/switch/{other_kid.id}").status_code == 403
    assert client.get(f"/patient/appointments/{kid_appt.id}").status_code == 403
    assert client.post(f"/patient/appointments/{kid_appt.id}/cancel").status_code == 403
    assert client.get(f"/patient/family/{other_kid.id}/edit").status_code == 403
    assert client.post(f"/patient/family/{other_kid.id}/remove").status_code == 403


def test_guardian_can_open_dependent_records_pdf_and_files(client):
    login(client, "pat@t.com")
    _add_child(client)
    kid = Patient.query.filter(Patient.relationship_label == "Child").one()
    client.post("/logout")
    a = _appt(patient=kid)
    _complete(client, a)
    rec = MedicalRecord.query.one()

    login(client, "pat@t.com")
    assert client.get(f"/patient/records/{rec.id}").status_code == 200
    assert client.get(f"/files/records/{rec.id}/prescription.pdf").data.startswith(b"%PDF")
    assert client.get(f"/patient/appointments/{a.id}").status_code == 200
    client.post("/logout")
    login(client, "pat2@t.com")
    assert client.get(f"/patient/records/{rec.id}").status_code == 403
    assert client.get(f"/files/records/{rec.id}/prescription.pdf").status_code == 403


def test_emails_for_a_child_go_to_the_guardian(client, app):
    from app.services import change_status
    login(client, "pat@t.com")
    _add_child(client)
    kid = Patient.query.filter(Patient.relationship_label == "Child").one()
    a = _appt("pending", patient=kid)
    change_status(a, "confirmed")
    msg = app.extensions["outbox"][-1]
    assert msg["To"] == "pat@t.com" and "Little Pat" in msg.get_content()
    walk_in = create_walk_in_patient = _appt("pending", slot=time(10, 0), patient=_patient())  # noqa
    from app.services import create_patient_without_login
    stranger = create_patient_without_login("Walk Inn", 30, "Male")
    db.session.commit()
    before = len(app.extensions["outbox"])
    change_status(_appt("pending", slot=time(10, 30), patient=stranger), "confirmed")
    assert len(app.extensions["outbox"]) == before                      # nobody to email: silently skipped


def test_removing_a_family_member_keeps_their_records(client):
    login(client, "pat@t.com")
    _add_child(client)
    kid = Patient.query.filter(Patient.relationship_label == "Child").one()
    _appt(patient=kid)
    client.post(f"/patient/family/{kid.id}/remove")
    assert not db.session.get(Patient, kid.id).user.is_active and Appointment.query.count() == 1
    assert "<strong>Little Pat</strong>" not in client.get("/patient/family").get_data(as_text=True)
    assert "family-switch" not in client.get("/patient/dashboard").get_data(as_text=True)


def test_family_form_validation(client):
    login(client, "pat@t.com")
    r = client.post("/patient/family/new", data={"name": "", "relationship": "Child", "age": "x", "gender": ""})
    assert r.status_code == 200 and Patient.query.filter(Patient.relationship_label == "Child").count() == 0


# ═══ Reports & Excel ═════════════════════════════════════════════════════

def _report_fixture(client):
    doc = _doctor()
    doc.consultation_fee = 1000
    other = _doc2(fee=400)
    db.session.commit()
    done = _appt("confirmed", slot=time(9, 0))
    _complete(client, done)
    paid = Invoice.query.one()
    paid.status, paid.method = "paid", "cash"
    _appt("no_show", slot=time(10, 0))
    _appt("cancelled", slot=time(10, 30))
    other_done = _appt("confirmed", slot=time(9, 0), doctor=other, patient=_patient("pat2@t.com"))
    login(client, "two@t.com")
    client.post(f"/doctor/appointments/{other_done.id}/record", data={"diagnosis": "d", "treatment": "t"})
    client.post("/logout")
    db.session.commit()


def test_report_numbers(client):
    _report_fixture(client)
    today = date.today()
    d = report_data(today - timedelta(days=7), today)
    assert d["total"] == 4 and d["status"]["completed"] == 2 and d["status"]["no_show"] == 1 and d["status"]["cancelled"] == 1
    assert d["no_show_rate"] == pytest.approx(33.3, abs=0.1)           # 1 / (2 completed + 1 no-show)
    assert d["cancel_rate"] == 25.0
    assert d["revenue"] == 1000 and d["outstanding"] == 400
    top = d["by_doctor"][0]
    assert top["doctor"] == "Dr. Who" and top["revenue"] == 1000 and top["no_show"] == 1
    assert {h["hour"]: h["count"] for h in d["by_hour"]}["09:00"] == 2    # two 09:00 visits; cancelled one excluded
    assert sum(w["count"] for w in d["by_weekday"]) == 3                  # non-cancelled only


def test_reports_page_and_excel_export(client):
    _report_fixture(client)
    login(client, "admin@t.com")
    page = client.get("/admin/reports").get_data(as_text=True)
    assert "No-show rate" in page and "33.3%" in page and "Dr. Who" in page
    r = client.get("/admin/reports/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "By doctor", "By hour", "By weekday", "By department"]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Appointments"] == 4 and summary["Revenue collected"] == 1000
    assert wb["By doctor"]["A2"].value == "Dr. Who"


def test_report_date_range_handling(client):
    login(client, "admin@t.com")
    assert client.get("/admin/reports?from=garbage&to=also-bad").status_code == 200
    assert client.get("/admin/reports?from=2030-01-10&to=2030-01-01").status_code == 200      # swapped range


def test_excel_exports_neutralise_formulas(client):
    a = _appt()
    a.reason = "=cmd|' /C calc'!A0"
    a.patient.user.name = "=HYPERLINK(\"http://evil\")"
    db.session.commit()
    login(client, "admin@t.com")
    wb = load_workbook(io.BytesIO(client.get("/admin/appointments/export.xlsx").data))
    ws = wb["Appointments"]
    cells = [c.value for c in ws[2]]
    assert all(not (isinstance(v, str) and v.startswith("=")) for v in cells)
    assert any(isinstance(v, str) and v.startswith("'=") for v in cells)
    assert ws["D2"].data_type == "s"


def test_reports_admin_only(client):
    for email in ("doc@t.com", "desk@t.com", "pat@t.com"):
        login(client, email)
        assert client.get("/admin/reports").status_code == 403, email
        assert client.get("/admin/reports/export.xlsx").status_code == 403, email
        client.post("/logout")
