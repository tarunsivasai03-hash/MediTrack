"""Admin reports: volumes, no-shows, revenue and busiest times, plus Excel export."""

import io
from collections import Counter, defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from app import db
from app.models import Appointment, Doctor, Invoice, User

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def _pct(part: int, whole: int) -> float:
    return round(part / whole * 100, 1) if whole else 0.0


def report_data(start: date, end: date) -> dict:
    """Everything the reports page and the Excel export show for appointments dated start..end."""
    A = Appointment
    in_range = (A.date >= start, A.date <= end)

    status = {s: 0 for s in ("pending", "confirmed", "completed", "cancelled", "no_show")}
    for st, n in db.session.query(A.status, func.count(A.id)).filter(*in_range).group_by(A.status):
        status[st] = n
    total = sum(status.values())
    resolved = status["completed"] + status["no_show"]

    # money, grouped per doctor and status
    money = defaultdict(lambda: {"paid": 0.0, "unpaid": 0.0})
    rows = (db.session.query(Invoice.doctor_id, Invoice.status, func.sum(Invoice.amount - Invoice.discount))
            .join(A, A.id == Invoice.appointment_id).filter(*in_range, Invoice.status.in_(("paid", "unpaid")))
            .group_by(Invoice.doctor_id, Invoice.status))
    for doc_id, st, amount in rows:
        money[doc_id][st] = float(amount or 0)

    per_doctor_status = defaultdict(Counter)
    for doc_id, st, n in db.session.query(A.doctor_id, A.status, func.count(A.id)).filter(*in_range).group_by(A.doctor_id, A.status):
        per_doctor_status[doc_id][st] = n

    doctors = Doctor.query.join(User, Doctor.user_id == User.id).order_by(User.name).all()
    by_doctor = []
    for d in doctors:
        c = per_doctor_status.get(d.id, Counter())
        n_total = sum(c.values())
        if not n_total and not money.get(d.id):
            continue
        res = c["completed"] + c["no_show"]
        by_doctor.append({
            "doctor": d.user.name, "department": d.department.name if d.department else d.specialization,
            "appointments": n_total, "completed": c["completed"], "no_show": c["no_show"],
            "no_show_rate": _pct(c["no_show"], res), "cancelled": c["cancelled"],
            "revenue": money[d.id]["paid"], "outstanding": money[d.id]["unpaid"],
        })
    by_doctor.sort(key=lambda r: -r["revenue"])

    # busiest times (non-cancelled bookings), bucketed in Python so it works on any database
    hours, weekdays = Counter(), Counter()
    for day, t in db.session.query(A.date, A.time).filter(*in_range, A.status != "cancelled"):
        hours[t.hour] += 1
        weekdays[day.weekday()] += 1
    hour_range = range(min(hours, default=8), max(hours, default=18) + 1)
    by_hour = [{"hour": f"{h:02d}:00", "count": hours.get(h, 0)} for h in hour_range]
    by_weekday = [{"day": DAYS[i], "count": weekdays.get(i, 0)} for i in range(7)]

    dept = defaultdict(lambda: {"appointments": 0, "revenue": 0.0})
    for r in by_doctor:
        dept[r["department"]]["appointments"] += r["appointments"]
        dept[r["department"]]["revenue"] += r["revenue"]
    by_department = sorted(({"department": k, **v} for k, v in dept.items()), key=lambda r: -r["revenue"])

    paid = sum(r["revenue"] for r in by_doctor)
    unpaid = sum(r["outstanding"] for r in by_doctor)
    return {
        "start": start, "end": end, "status": status, "total": total,
        "no_show_rate": _pct(status["no_show"], resolved), "cancel_rate": _pct(status["cancelled"], total),
        "revenue": paid, "outstanding": unpaid,
        "by_doctor": by_doctor, "by_hour": by_hour, "by_weekday": by_weekday, "by_department": by_department,
    }


# ── Excel ────────────────────────────────────────────────────────────────

_HEAD_FILL = PatternFill("solid", fgColor="1D4ED8")


def safe_cell(value):
    """Stop spreadsheet formula injection: text that starts with = + - @ is stored as plain text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], money_cols=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _HEAD_FILL
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([safe_cell(v) for v in r])
    for col in money_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = '#,##0.00'
    for i, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 3, 45)
    ws.freeze_panes = "A2"
    return ws


def report_workbook(data: dict) -> bytes:
    wb = Workbook()
    s = data["status"]
    _sheet(wb, "Summary", ["Metric", "Value"], [
        ["From", data["start"].isoformat()], ["To", data["end"].isoformat()],
        ["Appointments", data["total"]], ["Completed", s["completed"]], ["Cancelled", s["cancelled"]],
        ["No-shows", s["no_show"]], ["No-show rate (%)", data["no_show_rate"]],
        ["Cancellation rate (%)", data["cancel_rate"]],
        ["Revenue collected", data["revenue"]], ["Outstanding", data["outstanding"]],
    ], first=True)
    _sheet(wb, "By doctor",
           ["Doctor", "Department", "Appointments", "Completed", "No-shows", "No-show rate (%)", "Cancelled",
            "Revenue collected", "Outstanding"],
           [[r["doctor"], r["department"], r["appointments"], r["completed"], r["no_show"], r["no_show_rate"],
             r["cancelled"], r["revenue"], r["outstanding"]] for r in data["by_doctor"]], money_cols=(8, 9))
    _sheet(wb, "By hour", ["Hour", "Appointments"], [[r["hour"], r["count"]] for r in data["by_hour"]])
    _sheet(wb, "By weekday", ["Day", "Appointments"], [[r["day"], r["count"]] for r in data["by_weekday"]])
    _sheet(wb, "By department", ["Department", "Appointments", "Revenue collected"],
           [[r["department"], r["appointments"], r["revenue"]] for r in data["by_department"]], money_cols=(3,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def appointments_workbook(appointments) -> bytes:
    wb = Workbook()
    _sheet(wb, "Appointments",
           ["ID", "Date", "Time", "Patient", "Doctor", "Specialization", "Status", "Walk-in", "Reason"],
           [[a.id, a.date.isoformat(), a.time.strftime("%H:%M"), a.patient.user.name, a.doctor.user.name,
             a.doctor.specialization, a.status, "yes" if a.is_walk_in else "", a.reason or ""] for a in appointments],
           first=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
